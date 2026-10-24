"""A state's NEET PG fees as an Excel workbook a junior counsellor can look things up in (owner, 2026-10-24: "colleges
name upar se niche, upar courses ke naam, uske niche sub column mai seat type Gen / Management / NRI … figures lakhs
mai … hostel fees bhi … easy to use").

    node --env-file=.env.local scripts/poster/state_fees_doc.mjs --state Maharashtra --json <file.json>
    python scripts/poster/state_fees_xlsx.py <file.json> <out.xlsx>

Sheets: the MD/MS colleges (a row per college, a column group per branch with Gen / Mgmt / NRI under it, fees in
Rs lakh), the DNB / NBE hospitals (a row each: fee, seat types, branches), and how to read both. Same data and the
same caveat as the PDF: fees as listed for the counselling, never official, and the listing does not say for every
figure whether it is a year's fee or the whole course's. After writing, the workbook is read back and every fee
compared with the JSON; any difference stops the script.
"""

import json
import sys
from statistics import median

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SRC, OUT = sys.argv[1], sys.argv[2]
data = json.load(open(SRC, encoding="utf-8"))
STATE, AS_OF = data["state"], data["asOf"]

# Seat types per kind of college: which listed seat type fills the Gen, Mgmt and NRI columns.
SUB = ["Gen", "Mgmt", "NRI"]
COLUMN_OF = {
    "govt": {"Govt Quota": "Gen"},
    "private": {"Govt Quota": "Gen", "Management Quota": "Mgmt", "NRI Quota": "NRI"},
    "deemed": {"MNG": "Mgmt", "NRI": "NRI"},
}
TYPE_LABEL = {"govt": "Govt", "private": "Private", "deemed": "Deemed"}
SEAT_LABEL = {s["key"]: s["label"] for s in data["seats"]}

FONT = "Arial"
f = lambda **k: Font(name=FONT, **{"size": 10, **k})
fill = lambda hex_: PatternFill("solid", start_color=hex_, end_color=hex_)
NAVY, TEAL = "0B2447", "0F766E"
TYPE_FILL = {"govt": fill("DCFCE7"), "private": fill("E0F2FE"), "deemed": fill("FEF3C7")}
SUB_FILL = {"Gen": fill("D1FAE5"), "Mgmt": fill("DBEAFE"), "NRI": fill("FFEDD5")}
HEAD_FILLS = [fill("CCFBF1"), fill("E2E8F0")]
NOT_HERE = fill("EEF0F3")
FLAG = fill("FDE68A")
thin, med = Side(style="thin", color="CBD5E1"), Side(style="medium", color="64748B")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def lakh(n):
    """Rs -> lakh, two decimals, written the way it is said: 45, 1.57, 36.05."""
    v = round(n / 1e5, 2)
    return int(v) if v == int(v) else v


def rupees(n):
    s = str(int(round(n)))
    if len(s) <= 3:
        return "₹" + s
    head, tail = s[:-3], s[-3:]
    parts = []
    while len(head) > 2:
        parts.insert(0, head[-2:])
        head = head[:-2]
    return "₹" + ",".join(([head] if head else []) + parts + [tail])


def city_of(name):
    segs = [s.strip() for s in name.split(",")]
    return segs[-1] if len(segs) > 1 else ""


def pct50(xs):
    s = sorted(xs)
    return s[min(len(s) - 1, max(0, round(0.5 * (len(s) - 1))))]


def hostel_text(c):
    """One line for the hostel fee as listed. Seat types that list the same amount are one entry; a seat type that
    lists nothing defers to one that does; seat types that list different terms are named."""
    entries = []
    for seat, h in c["hostel"].items():
        h = {k: h.get(k) for k in ("min", "max", "period", "included", "mess")}
        for e in entries:
            same_amount = (h["min"], h["max"]) == (e["h"]["min"], e["h"]["max"]) or h["min"] is None or e["h"]["min"] is None
            fits = all(h[k] is None or e["h"][k] is None or h[k] == e["h"][k] for k in ("period", "included", "mess"))
            if same_amount and fits:
                e["h"] = {k: e["h"][k] if e["h"][k] is not None else h[k] for k in h}
                e["seats"].append(seat)
                break
        else:
            entries.append({"h": h, "seats": [seat]})
    entries = [e for e in entries if e["h"]["min"] is not None]

    def one(h):
        lo, hi = h["min"], h["max"] if h["max"] is not None else h["min"]
        amount = rupees(lo) if lo == hi else f"{rupees(lo)} – {rupees(hi)}"
        period = {"year": " a year", "month": " a month"}.get(h["period"] or "", " (period not listed)")
        extra = (", mess included" if h["mess"] else "") + (" — included in the fee" if h["included"] else "")
        return amount + period + extra

    if not entries:
        return "Not listed"
    if len(entries) == 1:
        return one(entries[0]["h"])
    return " · ".join(f"{' / '.join(SEAT_LABEL.get(s, s) for s in e['seats'])}: {one(e['h'])}" for e in entries)


def fee_groups(items):
    """[(fee or None, [course names])], cheapest first — for a seat type with no column of its own."""
    groups = {}
    for x in items:
        groups.setdefault(x["fee"], []).append(x["course"])
    return sorted(groups.items(), key=lambda g: (g[0] is None, g[0] or 0))


NAME = {c["course"]: c["name"] for c in data["courses"]}
matrix = [c for c in data["colleges"] if c["section"] in COLUMN_OF]
hospitals = [c for c in data["colleges"] if c["section"] == "nbe"]

# Every course any MD/MS college offers, in the shared order (MD, MS, diplomas; A–Z by the printed name).
offered = {x["course"] for c in matrix for items in c["seats"].values() for x in items}
courses = [c["course"] for c in data["courses"] if c["course"] in offered]
names = [NAME[c] for c in courses]
assert len(set(names)) == len(names), f"two courses print the same name: {sorted(n for n in names if names.count(n) > 1)}"

wb = Workbook()

# ------------------------------------------------------------------ sheet 1: MD/MS colleges
ws = wb.active
ws.title = "MD-MS colleges"
LEFT_COLS = [("S.No", 5), ("College", 46), ("Type", 8.5), ("City", 13), ("AIQ fee\n(MCC, every branch)", 11), ("Hostel fee (as listed)", 28)]
FIRST = len(LEFT_COLS) + 1  # first course column
NOTES_COL = FIRST + 3 * len(courses)
HEAD1, HEAD2 = 8, 9
TOP = HEAD2 + 1

ws["A1"] = f"{STATE} NEET PG 2026 — fees by college, branch and seat type"
ws["A1"].font = f(size=15, bold=True, color=NAVY)
ws["A2"] = "ALL FEES ARE IN ₹ LAKH  —  45 means ₹45,00,000  ·  1.57 means ₹1,57,000 (rounded to the nearest ₹1,000)"
ws["A2"].font = f(size=12, bold=True, color="B91C1C")
# The legend in short lines, so each fits a laptop screen beside the frozen columns.
LEGEND = [
    "Gen = State Govt Quota seat (CET Cell)   ·   Mgmt = Management Quota (CET Cell) / Paid seat at a deemed university (MCC)",
    "NRI = NRI Quota (CET Cell) / NRI seat at a deemed university (MCC)   ·   AIQ = All India Quota (MCC), government colleges, one fee for every branch",
    "Grey = branch not at this college   ·   – = not under this seat type   ·   NL = fee not listed   ·   Amber = far below the college's other fees, confirm   ·   Grey italic = 2025 fee",
    f"As listed for the 2026 NEET PG counselling (as of {AS_OF}) — not an official notice, and a year's fee or the whole course's is not always stated. Confirm with the college before paying.",
]
for i, line in enumerate(LEGEND, start=3):
    ws.cell(i, 1, line).font = f(size=10, color="334155") if i < 3 + len(LEGEND) - 1 else f(size=10, italic=True, color="92400E")

for i, (label, width) in enumerate(LEFT_COLS, start=1):
    ws.column_dimensions[get_column_letter(i)].width = width
    ws.cell(HEAD1, i).fill = fill(NAVY)
    cell = ws.cell(HEAD2, i, label)
    cell.font, cell.fill, cell.alignment = f(bold=True, color="FFFFFF"), fill(NAVY), CENTER

for k, course in enumerate(courses):
    col = FIRST + 3 * k
    ws.merge_cells(start_row=HEAD1, start_column=col, end_row=HEAD1, end_column=col + 2)
    head = ws.cell(HEAD1, col, NAME[course])
    head.font, head.alignment = f(bold=True, color=NAVY), CENTER
    for j, sub in enumerate(SUB):
        c = ws.cell(HEAD2, col + j, sub)
        c.font, c.fill, c.alignment = f(bold=True, size=9), SUB_FILL[sub], CENTER
        ws.cell(HEAD1, col + j).fill = HEAD_FILLS[k % 2]
        ws.column_dimensions[get_column_letter(col + j)].width = 7
        for r in (HEAD1, HEAD2):
            ws.cell(r, col + j).border = Border(left=med if j == 0 else thin, bottom=thin)
ws.cell(HEAD1, NOTES_COL).fill = fill(NAVY)
notes_head = ws.cell(HEAD2, NOTES_COL, "Notes (other seat types)")
notes_head.font, notes_head.fill, notes_head.alignment = f(bold=True, color="FFFFFF"), fill(NAVY), CENTER
ws.column_dimensions[get_column_letter(NOTES_COL)].width = 44
ws.row_dimensions[HEAD1].height = 30
ws.row_dimensions[HEAD2].height = 30
ws.cell(HEAD2, 5).comment = Comment("All India Quota seats at a government college, allotted by MCC. One fee for every branch.", "AdmissionHands")
for sub, text in {"Gen": "State Govt Quota seat, allotted by the CET Cell.",
                  "Mgmt": "Management Quota (CET Cell) at a private college; the Paid seat (MCC) at a deemed university.",
                  "NRI": "NRI Quota (CET Cell) at a private college; the NRI seat (MCC) at a deemed university."}.items():
    ws.cell(HEAD2, FIRST + SUB.index(sub)).comment = Comment(text, "AdmissionHands")

expected = {}  # (row, col) -> value, checked after the file is written back
for i, c in enumerate(matrix):
    r = TOP + i
    columns = COLUMN_OF[c["section"]]
    extra = [k for k in c["seats"] if k not in columns and k != "AIQ"]
    # One fee per course and column — the data holds one row per college, course and seat type.
    cell_of = {}
    for key, items in c["seats"].items():
        if key not in columns:
            continue
        fees = [x["fee"] for x in items if x["fee"] is not None]
        med_fee = pct50(fees) if len(fees) >= 3 else None
        for x in items:
            slot = (x["course"], columns[key])
            assert slot not in cell_of, f"two fees for {c['name']} {slot}"
            cell_of[slot] = {**x, "key": key, "low": med_fee is not None and x["fee"] is not None and x["fee"] < med_fee / 10}
    here = {x["course"] for items in c["seats"].values() for x in items}

    aiq = c["seats"].get("AIQ", [])
    aiq_fees = {x["fee"] for x in aiq}
    assert len(aiq_fees) <= 1, f"AIQ fee varies by branch at {c['name']}"
    left = [i + 1, c["name"], TYPE_LABEL[c["section"]], city_of(c["name"]),
            (lakh(next(iter(aiq_fees))) if next(iter(aiq_fees)) is not None else "NL") if aiq_fees else "–",
            hostel_text(c)]
    for j, v in enumerate(left, start=1):
        cell = ws.cell(r, j, v)
        cell.font = f(bold=(j == 2))
        cell.alignment = LEFT if j in (2, 4, 6) else CENTER
        cell.border = Border(bottom=thin)
    ws.cell(r, 3).fill = TYPE_FILL[c["section"]]
    if aiq_fees and next(iter(aiq_fees)) is not None:
        expected[(r, 5)] = lakh(next(iter(aiq_fees)))

    for k, course in enumerate(courses):
        for j, sub in enumerate(SUB):
            cell = ws.cell(r, FIRST + 3 * k + j)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(left=med if j == 0 else thin, bottom=thin)
            if course not in here:
                cell.fill = NOT_HERE
                continue
            x = cell_of.get((course, sub))
            if x is None:
                cell.value, cell.font = "–", f(color="94A3B8")
            elif x["fee"] is None:
                cell.value, cell.font = "NL", f(size=9, italic=True, color="64748B")
            else:
                cell.value = lakh(x["fee"])
                expected[(r, FIRST + 3 * k + j)] = cell.value
                old = x["session"] and str(x["session"]) != "2026"
                cell.font = f(italic=bool(old), color="6B7280" if old else "111827")
                if old:
                    cell.comment = Comment(f"{x['session']} fee — the 2026 fee is not listed yet.", "AdmissionHands")
                if x["low"]:
                    cell.fill = FLAG
                    cell.comment = Comment(f"{rupees(x['fee'])} — far below this college's other {SEAT_LABEL.get(x['key'], x['key'])} fees. "
                                           "Often a pre-clinical branch priced lower; confirm with the college.", "AdmissionHands")

    notes = []
    for key in extra:
        parts, groups = [], fee_groups(c["seats"][key])
        for fee, cs in groups:
            which = "every branch" if len(groups) == 1 else ", ".join(NAME[x] for x in sorted(cs, key=courses.index))
            parts.append(f"{lakh(fee) if fee is not None else 'not listed'} — {which}")
        notes.append(f"{SEAT_LABEL.get(key, key)} (₹ lakh): " + "; ".join(parts))
    note = ws.cell(r, NOTES_COL, "\n".join(notes) or None)
    note.alignment, note.font, note.border = LEFT, f(size=9), Border(bottom=thin, left=med)

    lines = max(1 + len(c["name"]) // 44, 1 + len(left[5]) // 27, sum(1 + len(n) // 42 for n in notes) if notes else 1)
    ws.row_dimensions[r].height = 14 * lines + 4

ws.freeze_panes = ws.cell(TOP, 4)  # header rows and S.No / College / Type stay in view
ws.auto_filter.ref = f"A{HEAD2}:{get_column_letter(NOTES_COL)}{TOP + len(matrix) - 1}"
ws.sheet_view.zoomScale = 90
ws.print_title_rows = f"{HEAD1}:{HEAD2}"
ws.print_title_cols = "A:C"
ws.page_setup.orientation = "landscape"
ws.page_setup.paperSize = ws.PAPERSIZE_A3

# ------------------------------------------------------------------ sheet 2: DNB / NBE hospitals
hs = wb.create_sheet("DNB & NBE hospitals")
hs["A1"] = f"{STATE} — DNB and NBE diploma hospitals (MCC counselling)"
hs["A1"].font = f(size=14, bold=True, color=NAVY)
hs["A2"] = "ALL FEES ARE IN ₹ LAKH  —  1.25 means ₹1,25,000"
hs["A2"].font = f(size=12, bold=True, color="B91C1C")
hs["A3"] = (f"As listed for the 2026 NEET PG counselling (as of {AS_OF}) — not an official notice; confirm with the MCC / NBEMS "
            "notice or the hospital. A hospital that also has MD/MS seats is on the first sheet too.")
hs["A3"].font = f(size=10, italic=True, color="92400E")
H_COLS = [("S.No", 5), ("Hospital", 44), ("Govt / Private", 10), ("City", 14), ("Fee (₹ lakh)", 10), ("Seat types", 22),
          ("DNB branches", 52), ("Diploma branches", 36), ("Hostel fee (as listed)", 28)]
for i, (label, width) in enumerate(H_COLS, start=1):
    hs.column_dimensions[get_column_letter(i)].width = width
    cell = hs.cell(5, i, label)
    cell.font, cell.fill, cell.alignment = f(bold=True, color="FFFFFF"), fill(NAVY), CENTER
hs.row_dimensions[5].height = 30
strip = lambda n: n.split(" ", 1)[1] if " " in n else n
for i, c in enumerate(sorted(hospitals, key=lambda c: c["name"])):
    r = 6 + i
    items = [x for v in c["seats"].values() for x in v]
    fees = sorted({x["fee"] for x in items if x["fee"] is not None})
    by_seat = {k: sorted({x["fee"] for x in v if x["fee"] is not None}) for k, v in c["seats"].items()}
    if len(fees) == 1:
        fee_cell = lakh(fees[0])
    elif fees:
        fee_cell = " · ".join(f"{SEAT_LABEL.get(k, k)} {', '.join(str(lakh(x)) for x in v)}" for k, v in by_seat.items() if v)
    else:
        fee_cell = "NL"
    dnb = sorted({NAME[x["course"]] for x in items if NAME[x["course"]].startswith("DNB")})
    dip = sorted({NAME[x["course"]] for x in items if not NAME[x["course"]].startswith("DNB")})
    row = [i + 1, c["name"], "Govt" if "Government" in (c["type"] or "") else "Private", city_of(c["name"]), fee_cell,
           ", ".join(SEAT_LABEL.get(k, k) for k in c["seats"]), ", ".join(strip(n) for n in dnb), ", ".join(strip(n) for n in dip),
           hostel_text(c)]
    for j, v in enumerate(row, start=1):
        cell = hs.cell(r, j, v)
        cell.font = f(bold=(j == 2), size=10 if j < 6 else 9)
        cell.alignment = CENTER if j in (1, 3, 5) else LEFT
        cell.border = Border(bottom=thin)
    if isinstance(fee_cell, (int, float)):
        expected[("h", r)] = fee_cell
    hs.cell(r, 3).fill = TYPE_FILL["govt" if row[2] == "Govt" else "private"]
    lines = max(1 + len(row[1]) // 42, 1 + len(row[6]) // 58, 1 + len(row[7]) // 40, 1 + len(row[8]) // 30)
    hs.row_dimensions[r].height = 13 * lines + 4
hs.freeze_panes = "C6"
hs.auto_filter.ref = f"A5:I{5 + len(hospitals)}"
hs.page_setup.orientation = "landscape"
hs.print_title_rows = "5:5"

# ------------------------------------------------------------------ sheet 3: how to read
rd = wb.create_sheet("How to read")
rd.column_dimensions["A"].width = 4
rd.column_dimensions["B"].width = 26
rd.column_dimensions["C"].width = 100
lines = [
    ("title", f"How to read this sheet — {STATE} NEET PG 2026 fees"),
    ("big", "All fees are in ₹ lakh. 45 means ₹45,00,000; 1.57 means ₹1,57,000 (rounded to the nearest ₹1,000)."),
    ("gap", ""),
    ("head", "Finding a fee"),
    ("step", ("1", "Find the college: press Ctrl+F and type part of its name, or use the filter arrows on Type (Govt / Private / Deemed) or City.")),
    ("step", ("2", f"Move right to the branch. Branch names are in row {HEAD1}; the college name stays on screen while you scroll.")),
    ("step", ("3", "Read the seat type under the branch: Gen, Mgmt or NRI. For a government college, the All India Quota fee is in the AIQ column (column E) — one fee for every branch.")),
    ("step", ("4", "Hostel fee is in column F, as the college lists it (a year, a month, or period not listed).")),
    ("step", ("5", "DNB and NBE diploma hospitals are on the second sheet: one fee (usually ₹1.25 lakh) and the branches each one has.")),
    ("gap", ""),
    ("head", "Seat types"),
    ("pair", ("Gen", "State Govt Quota seat, allotted by the State CET Cell. At a private college this is the state's fixed fee for that college.")),
    ("pair", ("Mgmt", "Management Quota at a private college (CET Cell). At a deemed university: the Paid seat, allotted by MCC.")),
    ("pair", ("NRI", "NRI Quota at a private college (CET Cell). At a deemed university: the NRI seat, allotted by MCC.")),
    ("pair", ("AIQ", "All India Quota seat at a government college, allotted by MCC.")),
    ("pair", ("Notes column", "Any other seat type a college lists (for example a Minority Quota), with its fee and branches.")),
    ("gap", ""),
    ("head", "What the cells mean"),
    ("pair", ("Grey cell", "The college does not have this branch.")),
    ("pair", ("–", "The college has the branch, but not under this seat type.")),
    ("pair", ("NL", "The seat is listed, but no fee is listed for it.")),
    ("pair", ("Amber cell", "Far below the same college's other fees for that seat type (often a pre-clinical branch priced lower). Hover for the exact amount; confirm with the college.")),
    ("pair", ("Grey italic", "Last year's (2025) fee — the 2026 fee is not listed yet. Hover to see it.")),
    ("gap", ""),
    ("head", "Before you quote a fee to a family"),
    ("text", f"These are the fees as listed for the 2026 NEET PG counselling (as of {AS_OF}). They are not an official notice."),
    ("text", "The listing does not say for every figure whether it is a year's fee or the whole course's — ask the college when it matters."),
    ("text", "Fees change. Confirm with the CET Cell or MCC notice, or with the college, before a student pays a fee or signs a bond."),
]
r = 1
for kind, v in lines:
    if kind == "title":
        rd.cell(r, 2, v).font = f(size=15, bold=True, color=NAVY)
    elif kind == "big":
        rd.cell(r, 2, v).font = f(size=12, bold=True, color="B91C1C")
    elif kind == "head":
        rd.cell(r, 2, v).font = f(size=12, bold=True, color=TEAL)
    elif kind == "step":
        rd.cell(r, 1, v[0]).font = f(bold=True, color=TEAL)
        rd.cell(r, 1).alignment = Alignment(horizontal="center", vertical="center")
        rd.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        rd.cell(r, 2, v[1]).alignment = LEFT
        rd.cell(r, 2).font = f()
        rd.row_dimensions[r].height = 28 if len(v[1]) > 110 else 16
    elif kind == "pair":
        a, b = rd.cell(r, 2, v[0]), rd.cell(r, 3, v[1])
        a.font, b.font, a.alignment, b.alignment = f(bold=True), f(), Alignment(vertical="center", indent=1), LEFT
        if v[0] in ("Gen", "Mgmt", "NRI"):
            a.fill = SUB_FILL[v[0]]
        elif v[0] == "Grey cell":
            a.fill = NOT_HERE
        elif v[0] == "Amber cell":
            a.fill = FLAG
        elif v[0] == "Grey italic":
            a.font = f(italic=True, color="6B7280")
        rd.row_dimensions[r].height = 28 if len(v[1]) > 100 else 16
    elif kind == "text":
        rd.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        rd.cell(r, 2, v).font, rd.cell(r, 2).alignment = f(), LEFT
    r += 1
rd.sheet_view.showGridLines = False

wb.save(OUT)

# ------------------------------------------------------------------ read it back: every fee where it should be
back = load_workbook(OUT)
m, h = back["MD-MS colleges"], back["DNB & NBE hospitals"]
bad = [(k, v, (h.cell(k[1], 5) if k[0] == "h" else m.cell(*k)).value) for k, v in expected.items()
       if (h.cell(k[1], 5).value if k[0] == "h" else m.cell(*k).value) != v]
assert not bad, f"{len(bad)} cells differ, e.g. {bad[:3]}"
numeric = sum(1 for k in expected if k[0] != "h")
listed = sum(1 for c in matrix for key, items in c["seats"].items() if key in COLUMN_OF[c["section"]] for x in items if x["fee"] is not None)
assert numeric - sum(1 for c in matrix if c["seats"].get("AIQ") and c["seats"]["AIQ"][0]["fee"] is not None) == listed, "fee count mismatch"
print(f"{OUT}\n  MD/MS sheet: {len(matrix)} colleges × {len(courses)} branches, {listed} fees + AIQ for "
      f"{sum(1 for c in matrix if c['seats'].get('AIQ'))} colleges · DNB/NBE sheet: {len(hospitals)} hospitals · read back: all match")
